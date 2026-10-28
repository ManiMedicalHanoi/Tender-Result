"""Gộp KQTT thô → 3 bộ dữ liệu không trùng: Nhãn khoa, Nha khoa, Ngoại khoa (Manipler).

Chạy:  python tools/kqtt_merge.py
Đọc:   raw/KQTT/<đợt>/*.xlsx, raw/KQTT/keywords.csv, raw/KQTT/product_groups.csv, rules/kqtt_rules.csv
Ghi:   processed/KQTT/  (xlsx + csv cho từng chuyên khoa, Loai_ra_tom_tat.csv, BAO_CAO.md;
       KQTT_Loai_ra.xlsx đầy đủ chỉ sinh ở máy, không đưa lên git)

Bước 1 · Khử trùng
  Hai dòng là MỘT khi khớp đủ 24 cột (mọi cột trừ STT) sau khi chuẩn hóa: Unicode NFC, gộp khoảng trắng,
  không phân biệt hoa thường. Dòng giống hệt nhau trong CÙNG một file là các mặt hàng khác nhau
  (vd. nhiều cỡ trâm cùng mô tả), nên số bản giữ lại = số lần xuất hiện nhiều nhất trong một file
  (không cộng dồn giữa các file).
Bước 2 · Phân loại (rules/kqtt_rules.csv), xét trên Tên thiết bị, vật tư y tế
  Một dòng thuộc nhóm X khi: tên khớp ít nhất 1 mẫu `ten` của X, và (nếu X có mẫu `ngu_canh`) tên, cấu hình,
  nhãn hiệu, hãng sản xuất hoặc tên CĐT khớp ít nhất 1 mẫu `ngu_canh`, và tên KHÔNG khớp mẫu `loai_tru` nào của X.
  Khớp nhiều nhóm → lấy nhóm đứng trước trong file quy tắc; các nhóm còn lại ghi ở cột `nhom_khop_khac`.
Bước 3 · Xuất
  Mỗi chuyên khoa một file gồm các dòng khớp nhóm sản phẩm của chuyên khoa đó.
  Dòng không khớp nhóm nào → KQTT_Loai_ra (để soát và bổ sung quy tắc).
"""
import collections
import csv
import datetime as dt
import os
import re
import sys
import unicodedata

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
RAW = os.path.join(BASE, "raw", "KQTT")
OUT = os.path.join(BASE, "processed", "KQTT")
RULES = os.path.join(BASE, "rules", "kqtt_rules.csv")

SPECIALTIES = [  # (chuyên khoa, tiền tố mã nhóm, tên file)
    ("Nhãn khoa", "MK-", "KQTT_Nhan_khoa"),
    ("Nha khoa", "NK-", "KQTT_Nha_khoa"),
    ("Ngoại khoa", "NG-MANIPLER", "KQTT_Ngoai_khoa_Manipler"),
]
EXTRA_COLS = ["chuyen_khoa", "nhom_sp_mmh", "ten_nhom", "nhom_khop_khac", "hang_mani",
              "tu_khoa_nguon", "so_file_nguon", "ban_sao_thu", "khoa_dong"]

# Vị trí dấu thanh kiểu cũ/mới (thuỷ/thủy, hoà/hòa) → một dạng, để quy tắc chỉ cần viết một kiểu
_TONE = {"uỷ": "ủy", "uý": "úy", "uỳ": "ùy", "uỹ": "ũy", "uỵ": "ụy",
         "oà": "òa", "oá": "óa", "oả": "ỏa", "oã": "õa", "oạ": "ọa",
         "oè": "òe", "oé": "óe", "oẻ": "ỏe", "oẽ": "õe", "oẹ": "ọe"}


def clean(v):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", "" if v is None else str(v))).strip()


def fold(s):
    s = clean(s).lower()
    for a, b in _TONE.items():
        s = s.replace(a, b)
    return s


def nfc(s):
    return unicodedata.normalize("NFC", s)


# ---------- quy tắc ----------
def load_rules():
    rules = collections.OrderedDict()
    with open(RULES, encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            g = rules.setdefault(r["ma_nhom"].strip(), {"ten": [], "ngu_canh": [], "loai_tru": []})
            g[r["loai"].strip()].append(re.compile(r["mau"], re.I))
    return rules


def classify(name, context, rules):
    """name: tên hàng; context: tên + cấu hình + nhãn hiệu + hãng + tên CĐT."""
    n, ctx = fold(name), fold(name) + " | " + fold(context)
    hits = []
    for code, g in rules.items():
        if not any(p.search(n) for p in g["ten"]):
            continue
        if g["ngu_canh"] and not any(p.search(ctx) for p in g["ngu_canh"]):
            continue
        if any(p.search(n) for p in g["loai_tru"]):
            continue
        hits.append(code)
    return hits


# ---------- đọc dữ liệu thô ----------
def load_raw():
    def read_csv(name):
        with open(os.path.join(RAW, name), encoding="utf-8-sig") as fh:
            return list(csv.DictReader(fh))

    kw_spec = {nfc(k["tu_khoa"]).lower(): k["chuyen_khoa"] for k in read_csv("keywords.csv")}
    groups = {g["ma_nhom"]: g for g in read_csv("product_groups.csv")}
    files = {nfc(r["file"]): r["tu_khoa"] for r in read_csv("manifest.csv")}

    header, per_key, first_row, raw_rows = None, collections.defaultdict(collections.Counter), {}, 0
    for batch in sorted(d for d in os.listdir(RAW) if re.fullmatch(r"\d{4}-\d{2}-\d{2}", d)):
        for f in sorted(os.listdir(os.path.join(RAW, batch))):
            if not f.lower().endswith(".xlsx"):
                continue
            kw = files.get(nfc(f))
            if kw is None:
                sys.exit(f"{f} chưa có trong manifest.csv: chạy tools/kqtt_manifest.py trước")
            rows = list(openpyxl.load_workbook(os.path.join(RAW, batch, f), read_only=True)
                        .worksheets[0].iter_rows(values_only=True))
            h = [clean(x) for x in rows[1]]
            if header is None:
                header = h
            elif h != header:
                sys.exit(f"{f}: tiêu đề cột khác các file khác")
            for r in rows[2:]:
                if not r or r[0] in (None, ""):
                    continue
                vals = [clean(x) for x in r]
                key = tuple(v.lower() for v in vals[1:])
                per_key[key][(batch, kw)] += 1
                first_row.setdefault(key, vals)
                raw_rows += 1
    return header, per_key, first_row, raw_rows, kw_spec, groups


# ---------- xuất ----------
def write_xlsx(path, header, rows, title):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(header)
    for r in rows:
        ws.append(r)
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")
    for c in ws[1]:
        c.font, c.fill = bold, fill
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"Tên thiết bị, vật tư y tế": 45, "Cấu hình, tính năng kỹ thuật cơ bản": 50, "Tên CĐT": 32,
              "Tên NT trúng thầu": 32, "Hãng sản xuất": 24, "ten_nhom": 26, "tu_khoa_nguon": 24, "khoa_dong": 14}
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
    wb.save(path)


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)


def main():
    rules = load_rules()
    header, per_key, first_row, raw_rows, kw_spec, groups = load_raw()
    unknown = [c for c in rules if c not in groups]
    if unknown:
        sys.exit(f"Mã nhóm trong quy tắc chưa có ở product_groups.csv: {unknown}")

    i_name = header.index("Tên thiết bị, vật tư y tế")
    i_spec = header.index("Cấu hình, tính năng kỹ thuật cơ bản")
    i_brand, i_maker = header.index("Nhãn hiệu"), header.index("Hãng sản xuất")
    i_ctx = [i_spec, i_brand, i_maker, header.index("Tên CĐT")]
    out_header = header[1:] + EXTRA_COLS

    buckets = {s[2]: [] for s in SPECIALTIES}
    rejected = []
    stats = collections.Counter()
    by_group = collections.Counter()
    overlap = collections.Counter()
    for n_key, (key, src) in enumerate(sorted(per_key.items(), key=lambda kv: first_row[kv[0]][i_name].lower()), 1):
        vals = first_row[key]
        copies = max(src.values())
        kws = sorted({kw for _, kw in src}, key=str.lower)
        hits = classify(vals[i_name], " | ".join(vals[i] for i in i_ctx), rules)
        mani = "Có" if re.search(r"\bmani\b", f"{vals[i_brand]} {vals[i_maker]}".lower()) else ""
        stats["unique_keys"] += 1
        stats["unique_rows"] += copies
        overlap[len(src)] += 1
        row_id = f"K{n_key:06d}"
        if hits:
            g = groups[hits[0]]
            spec_file = next(f for s, p, f in SPECIALTIES if hits[0].startswith(p))
            extra = [g["chuyen_khoa"], hits[0], g["ten_nhom"], "; ".join(hits[1:]), mani,
                     "; ".join(kws), len(src)]
            for c in range(1, copies + 1):
                buckets[spec_file].append(vals[1:] + extra + [c, row_id])
            by_group[hits[0]] += copies
        else:
            specs = sorted({sp.strip() for kw in kws for sp in kw_spec.get(nfc(kw).lower(), "").split(";")})
            extra = ["; ".join(s for s in specs if s), "", "", "", mani, "; ".join(kws), len(src)]
            for c in range(1, copies + 1):
                rejected.append(vals[1:] + extra + [c, row_id])
            stats["rejected_rows"] += copies
            if mani:
                stats["rejected_mani"] += copies

    os.makedirs(OUT, exist_ok=True)
    for spec, _, fname in SPECIALTIES:
        rows = buckets[fname]
        write_xlsx(os.path.join(OUT, fname + ".xlsx"), out_header, rows, spec)
        write_csv(os.path.join(OUT, fname + ".csv"), out_header, rows)
    # Bản đầy đủ (~18 MB) chỉ sinh ở máy, không đưa lên git (.gitignore); bản tóm tắt theo tên hàng thì commit
    write_xlsx(os.path.join(OUT, "KQTT_Loai_ra.xlsx"), out_header, rejected, "Loại ra")
    write_rejected_summary(rejected, out_header)

    write_report(stats, raw_rows, buckets, by_group, groups, overlap, rejected, out_header)
    print(f"Thô {raw_rows:,} → không trùng {stats['unique_rows']:,}; "
          + ", ".join(f"{s} {len(buckets[f]):,}" for s, _, f in SPECIALTIES)
          + f"; loại ra {stats['rejected_rows']:,} (trong đó MANI {stats['rejected_mani']})")


def write_rejected_summary(rejected, out_header):
    """Loai_ra_tom_tat.csv: mỗi dòng một tên hàng bị loại, kèm số dòng, từ khóa nguồn, có MANI không."""
    i_name, i_kw = out_header.index("Tên thiết bị, vật tư y tế"), out_header.index("tu_khoa_nguon")
    i_spec, i_mani = out_header.index("chuyen_khoa"), out_header.index("hang_mani")
    agg = {}
    for r in rejected:
        k = (r[i_name].lower(), r[i_kw])
        a = agg.setdefault(k, [r[i_name], r[i_kw], r[i_spec], 0, ""])
        a[3] += 1
        if r[i_mani]:
            a[4] = "Có"
    rows = sorted(agg.values(), key=lambda a: (-a[3], a[0].lower()))
    write_csv(os.path.join(OUT, "Loai_ra_tom_tat.csv"), ["ten_hang", "tu_khoa_nguon", "chuyen_khoa_tu_khoa", "so_dong", "hang_mani"], rows)


def write_report(stats, raw_rows, buckets, by_group, groups, overlap, rejected, out_header):
    i_kw = out_header.index("tu_khoa_nguon")
    rej_kw = collections.Counter()
    for r in rejected:
        for k in r[i_kw].split("; "):
            rej_kw[k] += 1
    lines = [
        "# Báo cáo gộp KQTT",
        "",
        f"_Sinh tự động bởi `tools/kqtt_merge.py` lúc {dt.datetime.now():%d/%m/%Y %H:%M}. Đừng sửa tay._",
        "",
        "## Khử trùng",
        f"- Dòng thô (tổng các file): **{raw_rows:,}**",
        f"- Dòng không trùng: **{stats['unique_rows']:,}** (bỏ {raw_rows - stats['unique_rows']:,} dòng lặp giữa các file)",
        "- Số file từ khóa chứa cùng một dòng: " + " · ".join(f"{k} file: {v:,}" for k, v in sorted(overlap.items())),
        "",
        "## Kết quả theo chuyên khoa",
        "| Chuyên khoa | File | Số dòng |",
        "|---|---|---|",
    ]
    for spec, _, f in SPECIALTIES:
        lines.append(f"| {spec} | `{f}.xlsx` / `.csv` | {len(buckets[f]):,} |")
    lines.append(f"| (không khớp nhóm nào) | `KQTT_Loai_ra.xlsx` (chỉ ở máy) · `Loai_ra_tom_tat.csv` | {stats['rejected_rows']:,} |")
    lines += ["", "## Theo nhóm sản phẩm MMH", "| Mã nhóm | Nhóm | Số dòng |", "|---|---|---|"]
    for code, g in groups.items():
        if by_group[code]:
            lines.append(f"| `{code}` | {g['ten_nhom']} | {by_group[code]:,} |")
    lines += ["", "## Dòng bị loại theo từ khóa nguồn (để soát quy tắc)", "| Từ khóa | Số dòng loại ra |", "|---|---|"]
    for k, v in rej_kw.most_common():
        lines.append(f"| {k} | {v:,} |")
    lines += ["", f"Dòng có nhãn hiệu/hãng MANI nhưng bị loại: **{stats['rejected_mani']}** "
              "(lọc cột `hang_mani` trong `Loai_ra_tom_tat.csv` để kiểm tra).", ""]
    with open(os.path.join(OUT, "BAO_CAO.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))


if __name__ == "__main__":
    main()
